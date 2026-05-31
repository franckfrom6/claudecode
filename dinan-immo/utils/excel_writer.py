# -*- coding: utf-8 -*-
"""
Gestion du fichier Excel des annonces.

Règles :
- Ne supprime JAMAIS les lignes existantes.
- Ajoute uniquement les annonces marquées "nouveau = True".
- Met en évidence les nouvelles lignes (fond jaune).
- Stocke un hash caché pour la déduplication entre sessions.
- Fige l'en-tête, largeurs de colonnes ajustées.
"""

import os
import logging
from datetime import datetime

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

from config import EXCEL_PATH, COLONNES_EXCEL

logger = logging.getLogger(__name__)

# Styles
_FILL_ENTETE = PatternFill(start_color="1F6AA5", end_color="1F6AA5", fill_type="solid")
_FILL_NOUVELLE = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
_FONT_ENTETE = Font(name="Calibri", color="FFFFFF", bold=True, size=11)
_FONT_LIEN = Font(name="Calibri", color="0563C1", underline="single", size=10)
_FONT_NORMALE = Font(name="Calibri", size=10)
_ALIGN_CENTRE = Alignment(horizontal="center", vertical="center")

# Largeurs de colonnes (en unités Excel)
_LARGEURS = {
    "Commune": 18,
    "Type": 13,
    "Loyer (€)": 12,
    "Charges": 10,
    "Chambres": 11,
    "Surface (m²)": 13,
    "DPE": 8,
    "Lien": 55,
    "Date repérée": 14,
    "Source": 20,
    "Nouveau": 10,
    "Score": 9,
    "Hash": 0,  # 0 = colonne masquée
}

# Colonnes du fichier (visibles + Hash caché)
_COLONNES_FICHIER = COLONNES_EXCEL + ["Hash"]


# ---------------------------------------------------------------------------
# Lecture
# ---------------------------------------------------------------------------

def charger_hashs_existants() -> set:
    """
    Lit le fichier Excel et retourne l'ensemble des hashs déjà présents.
    Retourne un set vide si le fichier n'existe pas encore.
    """
    if not os.path.exists(EXCEL_PATH):
        return set()

    try:
        wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
        ws = wb.active

        # Trouve la colonne "Hash" dans la première ligne
        col_hash = None
        for cell in ws[1]:
            if cell.value == "Hash":
                col_hash = cell.column
                break

        if col_hash is None:
            logger.warning("Colonne 'Hash' absente — impossible de charger les hashs existants")
            wb.close()
            return set()

        hashs = set()
        for row in ws.iter_rows(min_row=2, min_col=col_hash, max_col=col_hash, values_only=True):
            val = row[0]
            if val:
                hashs.add(str(val))

        wb.close()
        logger.info(f"Excel existant : {len(hashs)} hashs chargés")
        return hashs

    except Exception as e:
        logger.error(f"Erreur lecture Excel : {e}")
        return set()


# ---------------------------------------------------------------------------
# Écriture
# ---------------------------------------------------------------------------

def sauvegarder_annonces(annonces: list) -> str:
    """
    Ajoute les nouvelles annonces dans le fichier Excel.
    Crée le fichier s'il n'existe pas encore.
    Retourne le chemin du fichier.
    """
    os.makedirs(os.path.dirname(EXCEL_PATH), exist_ok=True)

    if os.path.exists(EXCEL_PATH):
        wb = openpyxl.load_workbook(EXCEL_PATH)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Annonces Dinan"
        _ecrire_entetes(ws)

    # Compteur des lignes ajoutées
    nb_ajoutees = 0
    for ann in annonces:
        if ann.get("nouveau"):
            _ecrire_ligne(ws, ann)
            nb_ajoutees += 1

    _ajuster_colonnes(ws)
    wb.save(EXCEL_PATH)

    logger.info(f"Excel sauvegardé : {nb_ajoutees} nouvelles lignes → {EXCEL_PATH}")
    return EXCEL_PATH


# ---------------------------------------------------------------------------
# Fonctions internes
# ---------------------------------------------------------------------------

def _ecrire_entetes(ws):
    """Écrit la ligne d'en-têtes avec mise en forme."""
    for col_idx, nom in enumerate(_COLONNES_FICHIER, start=1):
        cell = ws.cell(row=1, column=col_idx, value=nom)
        cell.fill = _FILL_ENTETE
        cell.font = _FONT_ENTETE
        cell.alignment = _ALIGN_CENTRE
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def _ecrire_ligne(ws, ann: dict):
    """Ajoute une annonce sur la prochaine ligne disponible."""
    row = ws.max_row + 1

    valeurs = [
        ann.get("commune") or "",
        ann.get("type") or "",
        ann.get("loyer"),
        ann.get("charges") or "",
        ann.get("chambres"),
        ann.get("surface"),
        ann.get("dpe") or "",
        ann.get("lien") or "",
        ann.get("date_reperee") or datetime.now().strftime("%Y-%m-%d"),
        ann.get("source") or "",
        "Oui" if ann.get("nouveau") else "Non",
        ann.get("score", 0),
        ann.get("hash") or "",
    ]

    col_lien_idx = _COLONNES_FICHIER.index("Lien") + 1

    for col_idx, valeur in enumerate(valeurs, start=1):
        cell = ws.cell(row=row, column=col_idx, value=valeur)
        cell.font = _FONT_NORMALE

        # Fond jaune sur les nouvelles annonces
        if ann.get("nouveau"):
            cell.fill = _FILL_NOUVELLE

        # Lien hypertexte cliquable
        if col_idx == col_lien_idx and valeur:
            cell.hyperlink = valeur
            cell.font = _FONT_LIEN


def _ajuster_colonnes(ws):
    """Applique les largeurs fixes et masque la colonne Hash."""
    for col_idx, nom in enumerate(_COLONNES_FICHIER, start=1):
        lettre = get_column_letter(col_idx)
        largeur = _LARGEURS.get(nom, 15)
        if largeur == 0:
            ws.column_dimensions[lettre].hidden = True
        else:
            ws.column_dimensions[lettre].width = largeur
